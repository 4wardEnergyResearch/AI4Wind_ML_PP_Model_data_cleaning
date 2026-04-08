# ai4wind_ml/preprocessing/power.py
"""Turbine power data loading, cleaning, and quality flagging."""
from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import scipy.ndimage as ndimage
import scipy.signal as signal

from ..config import TurbineConfig


# ---------------------------------------------------------------------------
# Helper: derive cut-in / cut-out / rated wind speeds from power curve
# ---------------------------------------------------------------------------

def _power_curve_specs(
    power_curve: pd.DataFrame, rated_power_kw: float
) -> tuple[float, float, float]:
    """Return (cut_in_ws, cut_out_ws, rated_ws) from a power curve DataFrame."""
    ws = power_curve["windspeed"]
    pw = power_curve["power"]
    cut_in = ws[pw > 0].iloc[0]
    cut_out = ws[(ws > 10) & (pw == 0)].iloc[0]
    rated = ws[pw >= rated_power_kw * 0.99].iloc[0]
    return float(cut_in), float(cut_out), float(rated)


class PowerDataPreprocessor:
    """Load raw turbine power data, apply multi-stage quality flagging,
    and export cleaning diagnostics.

    The flagging pipeline is ported from the original AI4WIND
    ``compute_power_and_flag.py`` and follows this order:

    1. Resample to regular 10-min intervals
    2. Clip to physical bounds
    3. Flag invalid data (flag 1)
    4. Flag by power-curve shift envelope (flag 6)
    5. Flag by sigma clipping (flag 5)
    6. Flag by quantile-based envelope (flag 2)

    Each stage only operates on rows that are still unflagged (flag == 0).
    """

    def __init__(self, config: TurbineConfig):
        self.config = config
        self._cleaning_info: pd.DataFrame | None = None

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def load_raw_data(self, path: str | Path | None = None) -> pd.DataFrame:
        """Load a raw turbine CSV into a DataFrame with a datetime index.

        Parameters
        ----------
        path : str or Path, optional
            Explicit path to the CSV.  If *None*, assumes
            ``data/turbine_data/raw/{turbine_id}.csv``.
        """
        if path is None:
            path = Path("data/turbine_data/raw") / f"{self.config.turbine_id}.csv"
        df = pd.read_csv(path, parse_dates=["timestamp"], index_col="timestamp")
        return df

    def clean(self, raw: pd.DataFrame) -> tuple[pd.Series, pd.DataFrame]:
        """Run the full cleaning pipeline on *raw* data.

        Returns
        -------
        power_series : pd.Series
            Cleaned power time series (only rows with flag == 0).
        cleaning_info : pd.DataFrame
            Full DataFrame including all flag columns and theoretical power.
        """
        cfg = self.config
        if cfg.power_curve is None:
            raise ValueError(
                "TurbineConfig.power_curve must be set for cleaning. "
                "Use TurbineConfig.from_database() or set it manually."
            )
        cut_in, cut_out, rated_ws = _power_curve_specs(
            cfg.power_curve, cfg.rated_power_kw
        )

        data = self._resample(raw)
        data = self._clip_outliers(data)
        data = _flag_invalid_data(data, cfg.rated_power_kw, cut_in)

        cc = cfg.cleaning_config
        if cc is None:
            from ..config import CleaningConfig
            cc = CleaningConfig()

        if cc.perform_pc_shift:
            data = _flag_pc_shift(
                data,
                power_curve=cfg.power_curve,
                shift_ws=cc.pc_shift_ws,
                shift_pw=cc.pc_shift_pw,
                rated_ws=rated_ws,
                cut_out_ws=cut_out,
                nominal_power=cfg.rated_power_kw,
                cut_threshold=cc.cut_thresh_pc,
            )
        if cc.perform_sigma_clipping:
            data = _flag_sigma(
                data,
                nominal_power=cfg.rated_power_kw,
                sc_mult=cc.sc_mult,
                cut_threshold=cc.cut_thresh_sc,
            )
        if cc.perform_quantile_mapping:
            data = _flag_quantiles(
                data,
                ws_cut_in=cut_in,
                ws_rated=rated_ws,
                ws_cut_out=cut_out,
                nominal_power=cfg.rated_power_kw,
                number_of_quantiles=cc.num_quantiles,
                threshold_fall=cc.thresh_fall,
                cut_threshold=cc.cut_thresh_qm,
            )

        # Theoretical power via power curve interpolation
        data["Power computed [kW]"] = np.interp(
            data["windspeed [m/s]"],
            cfg.power_curve["windspeed"],
            cfg.power_curve["power"],
        )

        self._cleaning_info = data
        power_clean = data.loc[data["Flag Computed"] == 0, "power [kW]"]
        return power_clean, data

    def get_target(self) -> pd.Series:
        """Full pipeline: load → clean → return cleaned power series."""
        raw = self.load_raw_data()
        power, _ = self.clean(raw)
        return power

    def get_metrics(self) -> dict[str, float]:
        """Return RMSE, MAE, BIAS (normalised by rated power) after cleaning."""
        if self._cleaning_info is None:
            raise RuntimeError("Call clean() before get_metrics().")
        data = self._cleaning_info
        valid = data[data["Flag Computed"] == 0]
        residual = valid["power [kW]"] - valid["Power computed [kW]"]
        nom = self.config.rated_power_kw
        return {
            "RMSE": float(np.sqrt((residual**2).mean()) / nom),
            "MAE": float(residual.abs().mean() / nom),
            "BIAS": float(residual.mean() / nom),
            "flagged_count": int((data["Flag Computed"] > 0).sum()),
        }

    def export_cleaning_report(
        self,
        output_dir: str | Path | None = None,
    ) -> Path:
        """Write a diagnostic scatter plot and summary CSV.

        Returns the directory where outputs were saved.
        """
        if self._cleaning_info is None:
            raise RuntimeError("Call clean() before export_cleaning_report().")

        if output_dir is None:
            output_dir = Path("data/turbine_data/cleaned/diagnostics")
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)

        data = self._cleaning_info
        tid = self.config.turbine_id
        pc = self.config.power_curve

        # --- Scatter plot with flag colours ---
        grouped = data.dropna(subset=["windspeed [m/s]", "power [kW]"]).groupby(
            "Flag Computed"
        )
        colors = {0: "C0", 1: "C1", 2: "C2", 3: "C5", 4: "C6", 5: "C3", 6: "C4"}
        labels = {
            0: "valid",
            1: "out of range",
            2: "quantile envelope",
            5: "sigma clipping",
            6: "power curve shift",
        }

        total = len(data.dropna(subset=["windspeed [m/s]", "power [kW]"]))
        counts = data["Flag Computed"].value_counts().to_dict()
        ratios = {k: v / total for k, v in counts.items()} if total else {}

        fig, (ax_pw, ax_hist) = plt.subplots(
            1, 2, sharey=True,
            gridspec_kw={"width_ratios": [4, 1]},
            figsize=(16, 9),
        )

        for flag, grp in grouped:
            label = labels.get(flag, f"flag {flag}")
            pct = ratios.get(flag, 0) * 100
            ax_pw.plot(
                grp["windspeed [m/s]"],
                grp["power [kW]"],
                ".",
                markersize=1,
                alpha=0.6,
                color=colors.get(flag, "grey"),
                label=f"{label} ({pct:.0f}%)",
            )

        if pc is not None:
            ax_pw.plot(
                pc["windspeed"], pc["power"],
                "k--", alpha=0.3, label="manufacturer power curve",
            )

        valid_ratio = ratios.get(0, 0) * 100
        ax_pw.set_title(
            f"Turbine {tid}: {100 - valid_ratio:.1f}% flagged"
        )
        ax_pw.set_xlabel("wind speed [m/s]")
        ax_pw.set_ylabel("power [kW]")
        ax_pw.legend(markerscale=10, fontsize="small", loc="upper left")
        ax_pw.grid(True)

        # Histogram of flagged data
        flagged_groups = [
            grp["power [kW]"] for flag, grp in grouped if flag != 0
        ]
        if flagged_groups:
            ax_hist.hist(
                flagged_groups,
                bins=max(1, round(self.config.rated_power_kw / 20)),
                density=True,
                orientation="horizontal",
                color=[colors.get(f, "grey") for f, _ in grouped if f != 0],
                stacked=True,
            )
        ax_hist.set_xlabel("log density")
        ax_hist.set_xscale("log")
        ax_hist.set_title("Flagged distribution")
        ax_hist.grid(True)

        fig.subplots_adjust(wspace=0)
        plt.tight_layout()
        fig.savefig(output_dir / f"{tid}.png", dpi=300)
        plt.close(fig)

        # --- Summary CSV ---
        summary = pd.DataFrame(
            {"count": counts, "ratio": ratios}
        ).sort_index()
        summary.index.name = "flag"
        summary.to_csv(output_dir / f"{tid}_summary.csv")

        return output_dir

    # ------------------------------------------------------------------
    # Private helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _resample(data: pd.DataFrame) -> pd.DataFrame:
        """Resample to regular 10-min intervals, taking the first value."""
        return data.resample("10min").first()

    @staticmethod
    def _clip_outliers(data: pd.DataFrame) -> pd.DataFrame:
        """Clip columns to physically sensible ranges."""
        if "power [kW]" in data.columns:
            data["power [kW]"] = data["power [kW]"].clip(-1000, 10000)
        if "windspeed [m/s]" in data.columns:
            data["windspeed [m/s]"] = data["windspeed [m/s]"].clip(0, 200)
        if "temperature [C]" in data.columns:
            data["temperature [C]"] = data["temperature [C]"].clip(-30, 50)
        if "wind direction [deg]" in data.columns:
            data["wind direction [deg]"] = data["wind direction [deg]"] % 360
        if "gondel position [deg]" in data.columns:
            data["gondel position [deg]"] = data["gondel position [deg]"] % 360
        return data


# ======================================================================
# Fleet-level cleaning report (xlsx)
# ======================================================================


def export_fleet_cleaning_report(
    results: list[dict],
    output_path: str | Path = Path("data/turbine_data/cleaned/diagnostics/Fleet Cleaning Report.xlsx"),
    plots_dir: str | Path | None = Path("data/turbine_data/cleaned/diagnostics"),
    suspicious_threshold: float = 10.0,
) -> Path:
    """Write a formatted Excel report summarising cleaning results across turbines.

    Parameters
    ----------
    results : list[dict]
        One dict per turbine with keys: ``turbine_id``, ``turbine_type``,
        ``total_rows``, ``flag_counts`` (dict mapping flag int to count),
        and optionally ``plot_path``.
    output_path : Path
        Where to write the xlsx file.
    plots_dir : Path or None
        Directory containing per-turbine diagnostic PNGs (for hyperlinks).
        If *None*, no hyperlinks are added.
    suspicious_threshold : float
        Percentage of flagged data above which a turbine is marked suspicious.

    Returns
    -------
    Path
        The path to the written xlsx file.
    """
    import openpyxl
    from openpyxl.formatting.rule import ColorScaleRule

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # --- Build DataFrame ---
    rows = []
    for r in results:
        total = r["total_rows"]
        fc = r.get("flag_counts", {})
        # Ensure integer keys
        fc = {int(k): int(v) for k, v in fc.items()}
        valid = fc.get(0, 0)
        flagged_pct = (1 - valid / total) * 100 if total else 0
        suspicious = flagged_pct > suspicious_threshold

        plot_path = None
        if plots_dir is not None:
            p = Path(plots_dir) / f"{r['turbine_id']}.png"
            if p.exists():
                plot_path = str(p.resolve())

        rows.append({
            "Turbine ID": r["turbine_id"],
            "Turbine type": r.get("turbine_type", ""),
            "Suspicious": suspicious,
            "Plot": plot_path,
            "Total dp.": total,
            "Ct. (valid)": fc.get(0, 0),
            "Rt. (valid)": fc.get(0, 0) / total if total else 0,
            "Ct. (rng)": fc.get(1, 0),
            "Rt. (rng)": fc.get(1, 0) / total if total else 0,
            "Ct. (quant)": fc.get(2, 0),
            "Rt. (quant)": fc.get(2, 0) / total if total else 0,
            "Ct. (sc)": fc.get(5, 0),
            "Rt. (sc)": fc.get(5, 0) / total if total else 0,
            "Ct. (pc)": fc.get(6, 0),
            "Rt. (pc)": fc.get(6, 0) / total if total else 0,
        })

    df = pd.DataFrame(rows)
    df.to_excel(output_path, index=False)

    # --- Format with openpyxl ---
    wb = openpyxl.load_workbook(output_path)
    ws = wb.active

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:{chr(64 + ws.max_column)}{ws.max_row}"

    # Column widths
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 12

    # Header: bold + left-aligned
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
        cell.alignment = openpyxl.styles.Alignment(horizontal="left")

    # Plot hyperlinks (column D)
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=4)
        if cell.value and cell.value != "Plot":
            link_path = str(cell.value)
            cell.value = f'=HYPERLINK("file:///{link_path}", "Plot")'
            cell.font = openpyxl.styles.Font(underline="single", color="0000FF")

    # Suspicious turbines in red (column C)
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=3)
        if cell.value is True:
            red = openpyxl.styles.Font(color="FF0000")
            cell.font = red
            ws.cell(row=row_idx, column=1).font = red  # Turbine ID
            ws.cell(row=row_idx, column=2).font = red  # Turbine type

    # Format ratio columns as percentages with color scales
    ratio_cols = {
        "G": "valid",   # Rt. (valid) — green=high is good
        "I": "flag",    # Rt. (rng)
        "K": "flag",    # Rt. (quant)
        "M": "flag",    # Rt. (sc)
        "O": "flag",    # Rt. (pc)
    }
    for col_letter, kind in ratio_cols.items():
        for row_idx in range(2, ws.max_row + 1):
            cell = ws[f"{col_letter}{row_idx}"]
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00%"

        col_range = f"{col_letter}2:{col_letter}{ws.max_row}"
        if kind == "valid":
            # High valid ratio = green
            rule = ColorScaleRule(
                start_type="min", start_color="FF0000",
                mid_type="percentile", mid_value=50, mid_color="FFFF00",
                end_type="max", end_color="00FF00",
            )
        else:
            # High flag ratio = red
            rule = ColorScaleRule(
                start_type="min", start_color="00FF00",
                mid_type="percentile", mid_value=50, mid_color="FFFF00",
                end_type="max", end_color="FF0000",
            )
        ws.conditional_formatting.add(col_range, rule)

    # Threshold note
    ws.cell(row=1, column=ws.max_column + 2).value = "Suspicious threshold:"
    ws.cell(row=1, column=ws.max_column).font = openpyxl.styles.Font(bold=True)
    ws.cell(row=2, column=ws.max_column).value = suspicious_threshold / 100
    ws.cell(row=2, column=ws.max_column).number_format = "0.00%"

    wb.save(output_path)
    wb.close()

    return output_path


# ======================================================================
# Flagging functions (module-level, ported from compute_power_and_flag.py)
# ======================================================================


def _flag_invalid_data(
    data: pd.DataFrame, nominal_power: float, ws_cut_in: float
) -> pd.DataFrame:
    """Flag 1: univariate bounds and NaN."""
    data["Flag Computed"] = 0
    ws = data["windspeed [m/s]"]
    pw = data["power [kW]"]

    # Wind speed out of range
    data.loc[(ws < 0) | (ws > 40), "Flag Computed"] = 1
    # Power out of range
    data.loc[(pw / nominal_power <= -0.05) | (pw / nominal_power > 1.05), "Flag Computed"] = 1
    # Zero power at high wind
    data.loc[(ws > ws_cut_in * 2) & (pw < 0.005 * nominal_power), "Flag Computed"] = 1
    # NaN in key columns
    data.loc[ws.isna() | pw.isna(), "Flag Computed"] = 1

    return data


def _flag_pc_shift(
    data: pd.DataFrame,
    *,
    power_curve: pd.DataFrame,
    shift_ws: float,
    shift_pw: float,
    rated_ws: float,
    cut_out_ws: float,
    nominal_power: float,
    cut_threshold: float = float("nan"),
) -> pd.DataFrame:
    """Flag 6: data outside a shifted power-curve envelope."""
    unflagged = data[data["Flag Computed"] == 0].copy()

    # Apply wind speed cut threshold
    if np.isfinite(cut_threshold):
        unflagged = unflagged[unflagged["windspeed [m/s]"] <= cut_threshold]

    pc = power_curve.copy()
    pc_ws = pc["windspeed"]
    pc_pw = pc["power"]

    # Shifted curves
    ws_right = pc_ws + shift_ws / 100 * rated_ws
    ws_left = pc_ws - shift_ws / 100 * rated_ws
    pw_up = pc_pw + shift_pw / 100 * nominal_power
    pw_down = pc_pw - shift_pw / 100 * nominal_power

    # Filter to valid region
    threshold = cut_out_ws - (shift_ws + 1) / 100 * rated_ws
    unflagged = unflagged[unflagged["windspeed [m/s]"] < threshold]

    ws_vals = unflagged["windspeed [m/s]"]
    pw_vals = unflagged["power [kW]"]

    # Above left-shifted curve → flag
    above = pw_vals > np.interp(ws_vals, ws_left, pw_up)
    # Below right-shifted curve → flag
    below = pw_vals < np.interp(ws_vals, ws_right, pw_down)

    flagged_idx = unflagged.index[above | below]
    data.loc[flagged_idx, "Flag Computed"] = 6
    return data


def _flag_sigma(
    data: pd.DataFrame,
    *,
    nominal_power: float,
    sc_mult: float = 3.0,
    cut_threshold: float = float("nan"),
) -> pd.DataFrame:
    """Flag 5: n-sigma clipping per power bin with Savitzky-Golay smoothing."""
    step = 1  # kW per bin
    sg_window_kw = nominal_power / 3
    sg_window = int(sg_window_kw / step)
    sg_poly = 3
    medfilt_window = int(sg_window / 5)
    if medfilt_window % 2 == 0:
        medfilt_window += 1
    medfilt_window = max(medfilt_window, 3)

    group_data = data[data["Flag Computed"] == 0]
    if np.isfinite(cut_threshold):
        group_data = group_data[group_data["windspeed [m/s]"] <= cut_threshold]

    pw_min = group_data["power [kW]"].min()
    pw_max = group_data["power [kW]"].max()
    if pd.isna(pw_min) or pd.isna(pw_max):
        return data

    bins = list(np.arange(pw_min - 0.5 * step, pw_max + 0.5 * step, step))
    if len(bins) < 2:
        return data
    grouped = group_data.groupby(pd.cut(group_data["power [kW]"], bins), observed=False)

    # Central values, medians, std devs per bin
    bins_cv = {}
    medians = {}
    std_devs = {}
    for key, df in grouped:
        bins_cv[key] = (key.left + key.right) / 2
        ws = df["windspeed [m/s]"].dropna()
        if len(ws) > 1:
            medians[key] = ws.median()
            std_devs[key] = ws.std()
        else:
            medians[key] = np.nan
            std_devs[key] = np.nan

    # Fill NaNs
    med_series = pd.Series(medians).ffill().bfill()
    std_series = pd.Series(std_devs).ffill().bfill()
    medians = med_series.to_dict()
    std_devs = std_series.to_dict()

    # Savitzky-Golay filter on medians
    med_vals = list(medians.values())
    if len(med_vals) < sg_window:
        sg_window = max(len(med_vals) // 2 * 2 - 1, sg_poly + 2)
    if sg_window < sg_poly + 2:
        return data
    try:
        med_sg = signal.savgol_filter(med_vals, sg_window, sg_poly)
    except Exception:
        med_sg = signal.savgol_filter(med_vals, sg_window, sg_poly, mode="nearest")
    medians_sg = dict(zip(medians.keys(), med_sg))

    # Compute intervals and apply median filter
    int_lower = {}
    int_upper = {}
    for key in grouped.groups:
        int_lower[key] = medians_sg[key] - sc_mult * std_devs[key]
        int_upper[key] = medians_sg[key] + sc_mult * std_devs[key]

    int_lower_mf = dict(zip(
        int_lower.keys(),
        ndimage.median_filter(list(int_lower.values()), size=medfilt_window, mode="constant"),
    ))
    int_upper_mf = dict(zip(
        int_upper.keys(),
        ndimage.median_filter(list(int_upper.values()), size=medfilt_window, mode="constant"),
    ))

    # Flag outliers
    for key, df in grouped:
        cv = bins_cv[key]
        if cv > nominal_power * 0.99:
            idx = df[df["windspeed [m/s]"] <= int_lower_mf[key]].index
        elif cv > nominal_power * 0.01:
            idx = df[
                (df["windspeed [m/s]"] <= int_lower_mf[key])
                | (df["windspeed [m/s]"] >= int_upper_mf[key])
            ].index
        else:
            idx = df[df["windspeed [m/s]"] >= int_upper_mf[key]].index
        data.loc[idx, "Flag Computed"] = 5

    return data


def _flag_quantiles(
    data: pd.DataFrame,
    *,
    ws_cut_in: float,
    ws_rated: float,
    ws_cut_out: float,
    nominal_power: float,
    number_of_quantiles: int = 100,
    threshold_fall: float = -1.0,
    cut_threshold: float = float("nan"),
) -> pd.DataFrame:
    """Flag 2: quantile-based lower envelope."""
    group_data = data[data["Flag Computed"] == 0]

    # Wind speed bins: fine below rated, coarser above
    vtest = list(np.arange(-0.05, ws_rated + 0.05, 0.1)) + list(
        np.arange(ws_rated + 0.15, ws_cut_out + 0.5, 0.5)
    )
    grouped = group_data.groupby(
        pd.cut(group_data["windspeed [m/s]"], vtest), observed=False
    )

    # Compute quantiles per wind speed bin
    quantiles = []
    index = []
    for key, df in grouped:
        pw = df["power [kW]"].dropna()
        if not pw.empty:
            index.append((key.left + key.right) / 2)
            quantiles.append(pw.quantile(np.linspace(0.0, 0.5, number_of_quantiles)))

    if not quantiles:
        return data

    df_q = pd.DataFrame(quantiles)
    df_q.index = index

    # Build mask: follow lowest quantile with positive derivative,
    # switch to next higher when it drops
    df_q["mask"] = -0.05 * nominal_power
    relevant = 0.0

    for col in df_q.columns:
        if col == "mask":
            continue
        deriv = df_q[col].diff()
        # Check slope in the cut-in region
        region = deriv[(df_q.index <= ws_cut_in * 2) & (df_q.index >= ws_cut_in)]
        if len(region) > 0 and np.mean(abs(region)) <= 1:
            continue
        if relevant > ws_rated * 1.5:
            break
        deriv = deriv[df_q.index >= relevant]
        falls = deriv <= threshold_fall
        if falls.sum() >= 1:
            first_fall = falls[falls].index[0]
            if first_fall - relevant >= 0:
                df_q.loc[relevant:, "mask"] = df_q.loc[relevant:, col].values
                relevant = first_fall

    # Apply flags
    for ws_val in df_q.index:
        max_power = df_q.loc[ws_val, "mask"]
        mask = (
            (data["windspeed [m/s]"] > ws_val)
            & (data["power [kW]"] < max_power)
            & (data["Flag Computed"] == 0)
        )
        if np.isfinite(cut_threshold):
            mask = mask & (data["windspeed [m/s]"] <= cut_threshold)
        data.loc[mask, "Flag Computed"] = 2

    return data
